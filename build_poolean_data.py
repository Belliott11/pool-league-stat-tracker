"""Builds poolean-external-data.js from the real Poolean site's own data export (an .xlsx someone
downloads from the site's export feature), one season at a time.

Each run saves that season's data to poolean-seasons/<YEAR>.json, then rebuilds
poolean-external-data.js from every season file in that folder, so importing a new season never
overwrites an old one. The export has no year in it (its README: one season per export), so the
season year is given on the command line and also used to date every party night.

Output globals (all `var`, so the app's season picker can swap them to another season):
  POOLEAN_SEASONS        { "2026": {rankings, record, together, against, cards, games, names,
                         awards}, ... }
  POOLEAN_SEASON_LIST    season years, oldest first
  POOLEAN_RANKINGS       every party night's power-ranking result (rank/percentile per player)
  POOLEAN_RECORD         each player's real overall win-loss across every game that season
  POOLEAN_TOGETHER       real pairwise win-loss as teammates, key "a|b" (slugs sorted)
  POOLEAN_AGAINST        real pairwise win-loss as opponents, key "a|b" (a's record facing b)
  POOLEAN_SEASON_CARDS   the site's own frozen end-of-season line per player (win%, power%, crowns,
                         best rank), used instead of recomputing so it always matches the site
  POOLEAN_GAMES          every real game (game number, date, both rosters, winner), sorted by
                         date then game number: the site's numbers alone aren't in play order
  POOLEAN_NAMES          slug -> display name, merged across every season
The single-season globals start out as the latest season's data.

Awards come from award_results (who won) and award_tally_long (the frozen point tally behind
it), when both sheets are present -- a season still being voted on has neither yet, so `awards`
is just missing until it closes; re-run this same import once it does and the results appear,
nothing to type by hand. This is the closed tally, not the raw per-voter ballots (those live in
award_ballots_long, deliberately never read here). app.js's own AWARD_LABELS/AWARD_STAT_KEYS/
AWARD_IS_DUO tables attach a label, closest tracked stat, and duo-or-not to each `key` here, so
this file only carries the two things it can't derive: the numbers and who they picked.

Usage: python build_poolean_data.py PATH_TO_EXPORT.xlsx --season YEAR
       python build_poolean_data.py --rebuild      (just re-combine the saved season files)
"""
import argparse
import json
from collections import defaultdict
from itertools import combinations
from pathlib import Path

HERE = Path(__file__).parent
SEASONS_DIR = HERE / "poolean-seasons"
OUT = HERE / "poolean-external-data.js"

MONTHS = {"Jan": "01", "Feb": "02", "Mar": "03", "Apr": "04", "May": "05", "Jun": "06",
          "Jul": "07", "Aug": "08", "Sep": "09", "Oct": "10", "Nov": "11", "Dec": "12"}


def party_date_to_iso(label, year):
    # "Sun Jun 7" -> "2026-06-07"
    _, mon, day = label.split()
    return f"{year}-{MONTHS[mon]}-{int(day):02d}"


def parse_export(path, year):
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)

    rankings_by_date = defaultdict(list)
    for party_no, date, slug, name, rank, field_size, pct in list(wb["rankings_long"].iter_rows(values_only=True))[1:]:
        rankings_by_date[date].append({"slug": slug, "rank": rank, "fieldSize": field_size, "pct": pct})
    rankings = [{"date": party_date_to_iso(d, year), "players": sorted(v, key=lambda p: p["rank"])}
                for d, v in sorted(rankings_by_date.items(), key=lambda kv: party_date_to_iso(kv[0], year))]

    record = defaultdict(lambda: {"w": 0, "l": 0})
    for game_no, date, slug, name, team, won in list(wb["game_players_long"].iter_rows(values_only=True))[1:]:
        record[slug]["w" if won else "l"] += 1
    record = {slug: {**wl, "gp": wl["w"] + wl["l"]} for slug, wl in record.items()}

    # TOGETHER is keyed by the sorted pair (same fact for both players); AGAINST is per ordered
    # pair (a's record facing b is b's mirror image, not the same number).
    together = defaultdict(lambda: {"w": 0, "l": 0})
    against = defaultdict(lambda: {"w": 0, "l": 0})
    games = []
    for game_no, gid, date, created, team_a, team_b, winner, a_size, b_size, live, events in list(wb["games"].iter_rows(values_only=True))[1:]:
        ta, tb = team_a.split("|"), team_b.split("|")
        a_won = winner == "A"
        for x, y in combinations(ta, 2):
            together["|".join(sorted((x, y)))]["w" if a_won else "l"] += 1
        for x, y in combinations(tb, 2):
            together["|".join(sorted((x, y)))]["w" if not a_won else "l"] += 1
        for x in ta:
            for y in tb:
                against[f"{x}|{y}"]["w" if a_won else "l"] += 1
                against[f"{y}|{x}"]["l" if a_won else "w"] += 1
        games.append({"n": game_no, "date": party_date_to_iso(date, year), "a": ta, "b": tb, "w": winner})
    together = {k: {**v, "gp": v["w"] + v["l"]} for k, v in together.items()}
    against = {k: {**v, "gp": v["w"] + v["l"]} for k, v in against.items()}
    games.sort(key=lambda g: (g["date"], g["n"]))

    cards, names = {}, {}
    if "season_cards" in wb.sheetnames:
        for slug, name, wins, losses, win_pct, power_pct, parties, crowns, best_rank in list(wb["season_cards"].iter_rows(values_only=True))[1:]:
            cards[slug] = {"w": wins, "l": losses, "winPct": win_pct, "powerPct": power_pct, "parties": parties, "crowns": crowns, "bestRank": best_rank}
            names[slug] = name
    if "players" in wb.sheetnames:
        for slug, name, *_ in list(wb["players"].iter_rows(values_only=True))[1:]:
            names.setdefault(slug, name)

    # Each award's frozen point tally (award_tally_long): borda_points for a single-candidate
    # award, pair_votes for the two duo awards (best-duo/worst-duo aren't ranked ballots, they're
    # just the most-named pair). Only the measure app.js's own votedStandings actually uses --
    # first_choice_votes is in the sheet too but nothing here reads it.
    awards = []
    if "award_results" in wb.sheetnames and "award_tally_long" in wb.sheetnames:
        DUO_KEYS = {"best-duo", "worst-duo"}
        tally = defaultdict(list)
        for award, measure, key, name, value in list(wb["award_tally_long"].iter_rows(values_only=True))[1:]:
            wanted = "pair_votes" if award in DUO_KEYS else "borda_points"
            if measure == wanted:
                tally[award].append({"slug": key, "name": name, "points": value})
        for award, winner_slugs, *_rest in list(wb["award_results"].iter_rows(values_only=True))[1:]:
            standings = sorted(tally.get(award, []), key=lambda r: -r["points"])
            awards.append({"key": award, "winners": winner_slugs.split("|"), "votedStandings": standings})

    return {"rankings": rankings, "record": record, "together": together, "against": against,
            "cards": cards, "games": games, "names": names, "awards": awards}


def rebuild():
    seasons = {}
    for f in sorted(SEASONS_DIR.glob("*.json")):
        seasons[f.stem] = json.loads(f.read_text(encoding="utf-8"))
    if not seasons:
        raise SystemExit(f"No season files in {SEASONS_DIR}; import one with --season first.")
    years = sorted(seasons, key=int)
    names = {}
    for y in years:
        names.update(seasons[y]["names"])
    for y in years:
        seasons[y]["names"] = names  # one name map for every season: a slug's name doesn't change
    latest = years[-1]
    compact = lambda v: json.dumps(v, separators=(",", ":"))
    OUT.write_text(
        "// Generated by build_poolean_data.py from the real Poolean site's own data exports,\n"
        "// one saved file per season in poolean-seasons/. See that script for what each global holds.\n"
        f"var POOLEAN_SEASONS = {compact(seasons)};\n"
        f"var POOLEAN_SEASON_LIST = {compact(years)};\n"
        f"var POOLEAN_RANKINGS = POOLEAN_SEASONS[\"{latest}\"].rankings;\n"
        f"var POOLEAN_RECORD = POOLEAN_SEASONS[\"{latest}\"].record;\n"
        f"var POOLEAN_TOGETHER = POOLEAN_SEASONS[\"{latest}\"].together;\n"
        f"var POOLEAN_AGAINST = POOLEAN_SEASONS[\"{latest}\"].against;\n"
        f"var POOLEAN_SEASON_CARDS = POOLEAN_SEASONS[\"{latest}\"].cards;\n"
        f"var POOLEAN_GAMES = POOLEAN_SEASONS[\"{latest}\"].games;\n"
        f"var POOLEAN_NAMES = POOLEAN_SEASONS[\"{latest}\"].names;\n",
        encoding="utf-8",
    )
    summary = ", ".join(f"{y}: {len(seasons[y]['rankings'])} parties / {len(seasons[y]['games'])} games" for y in years)
    print(f"{len(years)} season(s) -> {OUT} ({summary})")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("export", nargs="?", help="path to the site's .xlsx export")
    ap.add_argument("--season", type=int, help="the season year this export covers, e.g. 2026")
    ap.add_argument("--rebuild", action="store_true", help="only re-combine the saved season files")
    args = ap.parse_args()
    if not args.rebuild:
        if not args.export or not args.season:
            ap.error("give an export path and --season YEAR (or --rebuild)")
        SEASONS_DIR.mkdir(exist_ok=True)
        data = parse_export(args.export, args.season)
        dest = SEASONS_DIR / f"{args.season}.json"
        replaced = dest.exists()
        dest.write_text(json.dumps(data, separators=(",", ":")), encoding="utf-8")
        print(f"{'Replaced' if replaced else 'Saved'} season {args.season}: {len(data['rankings'])} party nights, "
              f"{len(data['games'])} games, {len(data['cards'])} season cards, "
              f"{len(data['awards'])} award{'s' if len(data['awards']) != 1 else ''} -> {dest}")
    rebuild()


if __name__ == "__main__":
    main()
